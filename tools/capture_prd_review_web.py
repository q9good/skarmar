"""Reproducible UI review screenshots; disposable localhost data only.

Run via tools/check-prd-web.mjs capture_prd_review_web.py. Core workflows use
the real UI; explicitly labelled historical fixtures use the authenticated API.
No production sessions, private files or remote application URLs are accepted.
"""
import calendar, json, os, re, subprocess
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import urlparse
from uuid import uuid4
from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect, sync_playwright
from prd_ui_helpers import reveal_control

BASE=os.environ['SKARMA_TEST_URL'].rstrip('/')
assert urlparse(BASE).hostname in {'localhost','127.0.0.1','::1'}
OUT=Path('.local/functional-review-2026-10-29');OUT.mkdir(parents=True,exist_ok=True)
shots=[];checks=[];errors=[]
source=subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()
with sync_playwright() as p:
    browser=p.chromium.launch(executable_path=os.environ.get('CHROMIUM_EXECUTABLE'),headless=True,args=['--no-sandbox'])
    ctx=browser.new_context(viewport={'width':390,'height':844},locale='zh-CN',accept_downloads=True)
    page=ctx.new_page();page.on('pageerror',lambda e:errors.append(str(e)))
    def btn(name,root=None):return (root or page).get_by_role('button',name=name,exact=True)
    def click(name):reveal_control(page,'button',name).click()
    def field(name):return reveal_control(page,'textbox',name)
    def select(name,value):reveal_control(page,'combobox',name).select_option(value)
    def check(name):reveal_control(page,'checkbox',name).click()
    def call(method,path,body=None):
        r=ctx.request.fetch(BASE+'/api'+path,method=method,data=body,headers={'X-Skarma-Request':'1','X-Skarma-Contract':'skarmar.prd.v1'})
        assert r.ok,(path,r.status,r.text());return r.json()
    def state():return call('GET','/prd/state')
    def close_all():
        for _ in range(8):
            if not page.get_by_role('dialog').count():return
            page.get_by_role('dialog').get_by_role('button',name=re.compile('^关闭')).first.click()
        assert not page.get_by_role('dialog').count()
    def shot(key,title,section,description,focus=None):
        if focus is not None:focus.evaluate("(el,block)=>el.scrollIntoView({block})",'start' if key in {'starter','import'} else 'center')
        else:page.evaluate("()=>{for(const e of document.querySelectorAll('*'))if(['auto','scroll'].includes(getComputedStyle(e).overflowY))e.scrollTop=0;}")
        page.wait_for_timeout(80)
        assert not errors,errors
        assert page.evaluate('document.documentElement.scrollWidth<=innerWidth+1')
        filename=f'{len(shots)+1:02d}-{key}.png';page.screenshot(path=str(OUT/filename))
        shots.append({'file':filename,'title':title,'section':section,'description':description,'width':page.viewport_size['width'],'height':page.viewport_size['height']})
        (OUT/'manifest.json').write_text(json.dumps({'source':source,'synthetic':True,'shots':shots,'checks':checks},ensure_ascii=False,indent=2)+'\n')
        print('CAPTURE '+filename+' '+title,flush=True)
    def entity(kind,value):
        id=str(uuid4());body={'id':id,'operation_id':str(uuid4()),'expected_version':None,'change_note':'','mutation':{'kind':kind,'value':value}}
        if kind=='method':body['method_create_mode']='full'
        return call('PUT','/prd/records/'+id,body)['value']
    def knowledge(kind,title,goal_ids,method_ids):
        id=str(uuid4());return call('POST','/prd/process',{'kind':'knowledge','value':{'id':id,'operation_id':str(uuid4()),'expected_version':None,'action':'create','input':{'id':id,'kind':kind,'content':'审核演示：'+title,'goal_ids':goal_ids,'method_ids':method_ids,'stage':'try' if kind=='experience' else None,'next_review_date':today if kind=='difficulty' else None},'details':{'title':title,'basis':'每次轮流机会','support':'示范一次'},'derived_from':[],'note':''}})
    page.goto(BASE);expect(field('账号')).to_be_visible()
    shot('login','登录与家庭设备选项','账号与初始设置','登录页保留受信任家庭设备和申请注册入口。图中没有填写密码。')
    click('申请注册');shot('registration','申请注册','账号与初始设置','管理员审批的注册表单；密码没有12位下限。注册、审批与停用另有完整认证测试。');click('返回登录')
    field('账号').fill('owner');page.get_by_label('密码',exact=True).fill('synthetic-prd-ui-password');click('登录');expect(btn('记录上课')).to_be_enabled()
    today=page.evaluate("new Intl.DateTimeFormat('sv-SE',{timeZone:'Asia/Shanghai',year:'numeric',month:'2-digit',day:'2-digit'}).format(new Date())")
    assert not state()['goals'] and not state()['lessons']
    shot('empty','首次进入空库','今天','没有自动创建正式目标、课程或训练记录；两个主入口保留。')
    click('我的账号');click('初始内容与参考');click('选取全部参考方法及关联方向');click('生成初始内容预览')
    expect(page.get_by_text('初始内容写入预览',exact=True)).to_be_visible()
    shot('starter','初始内容须先预览确认','账号与初始设置','演示明确选择六个长期方向、十种参考方法、六位老师和四门课程；预览前没有写入。',page.get_by_text('初始内容写入预览',exact=True))
    click('确认按预览创建');expect(page.get_by_text('已确认：新增 26 项，使用已有 0 项，跳过 0 项。',exact=True)).to_be_visible();close_all();click('返回')
    s=state();long=next(g for g in s['goals'] if g['input']['title']=='对人的兴趣、交往与互动')
    medium=entity('goal',{'title':'演示：一起参与轮流游戏','level':'medium','parent_id':long['id'],'meaning':'在共同游戏里关注对方的回应'})
    page.reload();expect(btn('记录上课')).to_be_enabled()
    click('目标');click('新建目标');field('本领名称').fill('演示：把球轮流给对方');click('补充目标说明（可选）')
    select('上级目标',medium['id']);select('观察方式','count');field('生活意义').fill('在喜欢的活动中保持来回互动');field('观察标准').fill('轮流机会中观察主动递球及需要的提醒');field('预期帮助条件（可选）').fill('成人先示范一次，再等待')
    shot('goal-form','目标录入与观察模板','目标与方法','名称和类型即可保存；意义、标准、帮助条件、上级和领域按需补充，示例不替代输入。',field('观察标准'))
    click('保存');expect(page.get_by_role('dialog')).to_have_count(0)
    short=next(g for g in state()['goals'] if g['input']['title']=='演示：把球轮流给对方')
    other=entity('goal',{'title':'演示：需要时表达求助','level':'short','criteria':'需要帮助时接受语言、手势或图片表达'})
    method=entity('method',{'title':'演示：先示范，再等待回应','steps':'跟随孩子的兴趣，示范一次后等待，按需要提供支持','goal_ids':[long['id'],short['id']],'observe':'是否愿意接回互动','scope':'熟悉且喜欢的游戏'})
    exp=knowledge('experience','演示：示范后留出等待时间',[short['id']],[method['id']])
    difficulty=knowledge('difficulty','演示：换材料时需要更多帮助',[short['id']],[method['id']])
    # Historical fixtures are explicit; no plan is treated as actual performance.
    def past_session(offset):
        id=str(uuid4());day=(date.fromisoformat(today)-timedelta(days=offset)).isoformat()
        return call('PUT','/prd/sessions/'+id,{'id':id,'operation_id':str(uuid4()),'expected_version':None,'action':'complete','input':{'business_date':day,'time_slot':'下午','plan':'审核演示历史计划'+('（专门删除）' if offset==0 else ''),'targets':[{'goal_id':short['id'],'role':'main'}],'focus_goal_ids':[long['id']],'summary':{'targets':[{'goal_id':short['id'],'performed':True,'observation':'审核演示：示范后参与了几次来回滚球'}],'confirmation':'used','methods':[{'id':method['id'],'version':1}],'tags':['游戏'],'signals':[]}},'change_note':''})
    for offset in [2,8,16,35]:past_session(offset)
    last=date.fromisoformat(today).replace(day=1)-timedelta(days=1)
    catalog=state()['catalog'];courses=[next(c for c in catalog if c['input']['name']==name) for name in ['认知','社交']];teachers=[next(c for c in catalog if c['input']['name']==name) for name in ['唐唐','彭彭']]
    for n,(c,t) in enumerate([(courses[0],teachers[0]),(courses[0],teachers[1]),(courses[1],teachers[0])]):
        id=str(uuid4());call('PUT','/prd/lessons/'+id,{'id':id,'operation_id':str(uuid4()),'expected_version':None,'action':{'kind':'observation','value':{'business_date':last.isoformat(),'course_id':c['id'],'teacher_id':t['id'],'notes':'审核演示：上月正式课记'+str(n+1)}}})
    page.reload();expect(btn('记录上课')).to_be_enabled();shot('today','长期关注与当天入口','今天','前三个长期方向逐项分隔，其余方向可展开；今天只列当天内容。')
    page.set_viewport_size({'width':360,'height':640});expect(btn('记录上课')).to_be_in_viewport(ratio=1);expect(btn('记录一次家庭干预')).to_be_in_viewport(ratio=1)
    shot('today-small','小屏首屏仍可记课或计划','今天','360×640桌面手机尺寸下，六个真实长度的参考方向不会挤走两个主按钮；不是手机真机截图。');page.set_viewport_size({'width':390,'height':844})
    click('查看长期关注');shot('long-focus','长期方向详情与关联','今天','显示意义、按明确目标筛选的最近观察，以及关联的方法、经验和困难。');close_all()
    click('目标');shot('goals','目标目录','目标与方法','首屏直接看到目标及覆盖；搜索和期限、来源等筛选默认折叠。')
    page.set_viewport_size({'width':1440,'height':900});shot('goals-desktop','电脑宽屏目标目录','目标与方法','同一组数据在1440×900宽屏检查；内容限宽，避免输入和卡片无限拉伸。');page.set_viewport_size({'width':390,'height':844})
    click('查看目标：'+short['input']['title']);shot('goal-links','目标详情与方法、经验、困难','目标与方法','核对目标下能看到已关联方法、经验和困难，并可打开原记录和复盘。')
    close_all();click('方法');expect(btn('查看方法：'+method['input']['title'])).to_be_visible();shot('methods','方法按近况排列','目标与方法','方法显示关联目标和排序原因；困难优先提示检查与调整。')
    click('编辑方法：'+method['input']['title']);field('观察意图／可选').fill('审核演示：留意孩子是否愿意接回互动');field('修改说明').fill('审核演示：补充观察意图')
    shot('method-edit','方法编辑保留目标关联','目标与方法','方法名称、具体做法与目标关系可修改，做法调整及文字纠错区分处理。',field('具体做法／必填'))
    click('保存');expect(page.get_by_role('dialog')).to_have_count(0);method=next(m for m in state()['methods'] if m['id']==method['id'])
    click('今天');click('记录一次家庭干预');click('选择本次目标');check('选择目标：'+short['input']['title']);click('完成目标选择');field('准备怎么做').fill('审核演示：一起滚球，等待孩子回应，再轮到对方')
    shot('plan','干预计划与重点提醒','家庭干预','选中目标后显示相关困难和经验；主要练、顺便练和复习可调整，计划不计实际次数。')
    click('查看全部方向的近况');select('近况展示方式','ring');expect(page.get_by_role('img',name=re.compile('当前显示项覆盖环图'))).to_be_visible()
    shot('balance','计划中的均衡覆盖提示','家庭干预','目标、实际方法、活动可切换；来源分家庭、机构和合计。环图表示覆盖数量，不表示能力百分比。',page.get_by_role('img',name=re.compile('当前显示项覆盖环图')))
    click('关闭全部方向的近况');click('保存共享草稿并退出');expect(page.get_by_role('dialog')).to_have_count(0);expect(btn('填写今天的总结')).to_be_visible()
    planned=next(s for s in state()['sessions'] if s['status']=='planned');assert len(state()['sessions'])==5
    click('填写今天的总结');field('实际表现：'+short['input']['title']).fill('审核演示：示范后主动递球两次，换材料时需要休息')
    select('实际方法确认','used');check('实际方法：'+method['input']['title'])
    shot('summary','训练后填写实际表现与方法','家庭干预','实际表现先填；实际方法必须明确确认，不把计划方法自动当成用过。')
    click('新增本次困难');field('实际情况／必填 1').fill('审核演示：换成新球时离开活动');field('成人回应／可选 1').fill('暂停、陪伴休息，再提供选择')
    shot('summary-difficulty','总结沉淀新困难','家庭干预','实际情况、成人回应、原因猜测和条件分开；新内容默认关联本次明确目标，可调整。',field('实际情况／必填 1'))
    click('正式提交干预总结');expect(page.get_by_role('dialog')).to_have_count(0)
    saved=next(s for s in state()['sessions'] if s['id']==planned['id']);assert saved['status']=='completed' and saved['actual_methods'][0]['id']==method['id']
    checks.append('计划→共享草稿→实际总结使用同一ID；目标/实际方法/困难写入已核对，计划阶段不计实际覆盖')
    shot('today-saved','今天的已同步记录','今天','总结提交后当天列表显示作者与时间；历史记录不混入今天。')
    click('记录上课');click('课程名称（必填）：请选择课程')
    shot('course-picker','课程下拉、新增与历史排序','机构课程','已有选项可选、可搜索新增；选课程后老师按组合历史频次排序，不自动替换另一个字段。')
    click('认知');click('授课老师（必填）：请选择老师');click('唐唐');shot('lesson-basic','老人快速记课','机构课程','默认当天，课程与老师必填；选完可以直接提交，观察不是必填。')
    click('提交记录');expect(page.get_by_role('dialog')).to_have_count(0);expect(btn('补充课堂观察')).to_be_visible();shot('lesson-saved','记课后的首屏确认','机构课程','服务器回执后才显示已记下，补充观察入口在前面，补充不新增课时。')
    lesson=next(l for l in state()['lessons'] if l['observation']['business_date']==today)
    click('补充课堂观察');field('今天课上怎么样？').fill('审核演示：轮流传球，老师等待孩子回应后再改变材料')
    click('老师的做法和建议');field('老师的做法和建议').fill('审核演示：回家可在熟悉游戏里练习轮流和求助');check('请爸爸妈妈看看')
    shot('lesson-observation','可选课堂观察与家人待办','机构课程','自由观察先填，主题按需展开；老师建议不会自动变成下一步，须由家人明确整理。',field('老师的做法和建议'))
    click('提交记录');expect(page.get_by_role('dialog')).to_have_count(0);assert len(state()['lessons'])==4
    click('完成');btn(re.compile('^查看课记：')).first.click();click('整理这条记录');check(short['input']['title']);check('实际方法：'+method['input']['title'])
    shot('lesson-links','课堂记录关联目标与实际方法','机构课程','整理和原观察独立保存版本；机构覆盖与家庭覆盖分列，不能把方法适用关系当成实际使用。')
    select('这条记录如何处理','complete');click('保存整理');expect(page.get_by_role('dialog')).to_have_count(0)
    organized=next(l for l in state()['lessons'] if l['id']==lesson['id']);assert organized['organization']['targets'][0]['goal']['id']==short['id']
    click('完成');click('记录');shot('classroom','机构课堂记录','记录与统计','当前自然月按课程分组，保留原课堂观察和老师建议，可打开原记录。')
    click('课时核对');expect(btn('唐唐：2 节')).to_be_visible();shot('teacher-totals','上月按老师核对课时','记录与统计','默认上个自然月、按老师汇总节数，点击分组看明细；金额线下算。')
    click('按课程汇总');expect(btn('认知：2 节')).to_be_visible();click('认知：2 节');shot('course-totals','同一时期按课程汇总及明细','记录与统计','同一批三条上月记录，按课程计认知2节、社交1节；没有把节数换算为分钟。')
    click('家庭干预');shot('calendar','家庭干预月历','记录与统计','向下滚动后的完整月份区域；按月展示已完成次数，待总结与完成分开，空白表示未记录，瞬间不计入。',btn(today+' · 已完成1次').locator('..'))
    click('按目标查看时间分布');shot('history','目标的频次与时间远近','记录与统计','手机用13周小条和最近日期；可展开近90天或最近30次，历史层级和方法版本保留。');close_all()
    call('POST','/auth/users',{'username':'review-family','display_name':'审核演示家人','password':'s','role':'member'})
    admin_page=page;member_ctx=browser.new_context(viewport={'width':390,'height':844},locale='zh-CN');page=member_ctx.new_page();page.on('pageerror',lambda e:errors.append(str(e)))
    page.goto(BASE);field('账号').fill('review-family');page.get_by_label('密码',exact=True).fill('s');click('登录');expect(btn('记录上课')).to_be_enabled();click('记录');click('家庭干预');btn(re.compile(r'^\d{4}-\d{2}-\d{2} · .+ · 已完成 ·')).first.click()
    expect(btn('修改已完成总结')).to_be_visible();expect(btn('删除干预（管理员）')).to_have_count(0)
    shot('member','普通家人可修改总结，不能删除','观察与可靠性','用独立成员账号实际登录；保留修改已完成总结，管理员删除入口不出现。API权限拒绝另由完整测试验证。',btn('修改已完成总结'));member_ctx.close();page=admin_page
    disposable=past_session(0);page.reload();expect(btn('记录上课')).to_be_enabled();click('记录');click('家庭干预');btn(re.compile('.*已完成.*审核演示历史计划（专门删除）.*')).click();click('删除干预（管理员）')
    shot('admin-delete','管理员删除前明确确认','观察与可靠性','本次真实删除一条专门生成的合成干预；删除不连带清除沉淀内容，原来的示例总结和机构记录保持原样。')
    click('确认干预操作');expect(page.get_by_role('dialog')).to_have_count(0);assert next(s for s in state()['sessions'] if s['id']==disposable['id'])['status']=='deleted';checks.append('独立普通成员无删除入口；管理员实际删除专用合成记录并核对状态')
    click('跟进');shot('follow-tasks','跟进待处理事项','跟进与复盘','待办、课堂请求和到期困难集中显示；这个分段不出现新建经验／困难入口。')
    click('经验库与困难');shot('library','经验库与困难','跟进与复盘','新建入口在此分段；经验阶段、困难状态、条件与来源分别保留。')
    click('新建经验');field('经验名称（必填）').fill('演示：从等待到主动回应');field('做法与依据').fill('审核演示：在喜欢的活动里等待回应，换材料时保留支持');shot('new-experience','独立录入经验','跟进与复盘','经验可记录已有做法或待验证想法，阶段由家人明确选择；名称、依据和关联不会通过示例自动填写。');click('选择关联目标与方法')
    check('关联目标：'+long['input']['title']);click('展开关联层级：'+long['input']['title']);click('展开关联层级：'+medium['input']['title']);check('关联目标：'+short['input']['title']);click('展开关联层级：'+short['input']['title'])
    shared=page.get_by_role('checkbox',name='关联方法：'+method['input']['title'],exact=True);expect(shared).to_have_count(2);shared.first.click();expect(shared.last).to_be_checked()
    shot('hierarchy','父子目标与方法分层多选','跟进与复盘','长中短目标按真实父子关系展开；同一方法出现在多个目标下，勾选状态共用，不自动勾选上下级。')
    close_all();click('经验库与困难');click('查看知识：'+exp['details']['title']);click('记录本次经验反馈');select('本次结果','helpful');field('本次实际观察').fill('审核演示：熟悉材料中愿意再来一轮')
    shot('experience-feedback','经验的实际反馈','跟进与复盘','有帮助、不确定、没帮上同等可选；每次保留条件和来源，其他家人可补看法，不靠点赞判断有效。');click('保存跟进');expect(page.get_by_role('dialog',name='记录反馈与跟进',exact=True)).to_have_count(0);close_all()
    click('查看知识：'+difficulty['details']['title']);click('记录本次困难跟进');select('本次结果','same');field('本次实际观察').fill('审核演示：换材料后仍需要提示，回到熟悉材料后继续')
    shot('difficulty','困难跟进与可比条件','跟进与复盘','记下尝试、成人回应和实际变化；原因猜测不当成事实，可比观察不足不自动升优先级。');click('保存跟进');close_all()
    click('目标');click('查看目标：'+short['input']['title']);click('现在复盘这个对象');field('复盘依据与说明').fill('审核演示：熟悉材料时主动回应增加，换材料仍需要支持')
    select('本次复盘决定','adjust');field('调整后的标准或做法（可选）').fill('先稳定熟悉活动，再只改变一种材料');click('新增复盘下一步');field('复盘下一步 1').fill('演示：请老师示范一次换材料的支持')
    shot('review','复盘实际依据与下一步','跟进与复盘','逐项回看已读取的记录，明确继续、调整或暂停；保存后才生成下一次安排，未读依据不会被自动消耗。')
    shot('review-decision','复盘决定与调整说明','跟进与复盘','看过依据后再决定继续、调整或暂停，说明与下一步需明确保存；可选日程和家人负担不会阻碍记录。',page.get_by_role('combobox',name='本次复盘决定',exact=True))
    click('保存复盘');expect(page.get_by_role('dialog',name='逐项复盘',exact=True)).to_have_count(0);close_all();click('跟进');click('维护下一步：演示：请老师示范一次换材料的支持')
    shot('action','下一步、负责人及状态','跟进与复盘','下一步可指派、排日期、纳入共享计划；已经在计划中的行动受占用保护。');close_all()
    click('目标');click('查看目标：'+short['input']['title']);click('查看观察与实际测量');click('新增观察指标');field('指标名称').fill('演示：轮流回应机会');field('预期帮助条件（可选）').fill('示范一次');click('保存观察');click('记录一次测量：演示：轮流回应机会');field('机会数').fill('5');field('做到数').fill('3');field('实际观察说明').fill('审核演示：同样材料、示范一次');click('保存观察')
    expect(page.get_by_text('做到 3／5 次，60%（该次机会内的比例）',exact=True)).to_be_visible();shot('measurement','实际测量与口径分段','观察与可靠性','3/5仅表示这次机会内的比例；记录次数不是能力进展，标准或帮助条件改变后分开看。',page.get_by_text('做到 3／5 次，60%（该次机会内的比例）',exact=True));close_all()
    click('今天');click('记一个瞬间');field('想记下的瞬间').fill('审核演示：拿着新球邀请家人一起玩');
    if not btn('关联：'+short['input']['title']).count():click(re.compile('^更多观察目标'))
    click('关联：'+short['input']['title']);shot('moment','随手记一个瞬间','观察与可靠性','单文本配日期与可选目标；正式保存可作观察依据，但不增加家庭干预次数。');click('保存观察');expect(page.get_by_role('dialog')).to_have_count(0)
    # A real concurrent change must preserve current input and show both versions.
    click('目标');click('编辑目标：'+short['input']['title']);field('生活意义').fill('审核演示：本窗口尚未提交的说明')
    current=next(g for g in state()['goals'] if g['id']==short['id']);call('PUT','/prd/records/'+current['id'],{'id':current['id'],'operation_id':str(uuid4()),'expected_version':current['version'],'change_note':'审核演示：另一位家人修改','mutation':{'kind':'goal','value':{**current['input'],'meaning':'审核演示：另一位家人已保存的说明'}}})
    page.evaluate("window.dispatchEvent(new Event('focus'))");expect(page.get_by_text('另一位家人已修改这条记录，你的填写仍保留。',exact=True)).to_be_visible();expect(btn('保存')).to_be_disabled()
    shot('conflict','多人修改冲突保护','观察与可靠性','服务器新版本和本窗口输入并列保留，禁止静默覆盖；须明确采用最新或合并。');close_all()
    click('今天');ctx.set_offline(True);click('记录上课');click('课程名称（必填）：请选择课程');click('社交');click('授课老师（必填）：请选择老师');click('彭彭');click('提交记录');expect(page.get_by_role('dialog')).to_have_count(0);expect(page.get_by_text(re.compile('^待上传课记与选项：'))).to_be_visible()
    shot('offline','离线仅保存在本机，联网补传','观察与可靠性','待上传单独提示且不计正式课时；联网后核对原账号、目录依赖及同一次请求，再显示服务器确认。')
    count=len(state()['lessons']);ctx.set_offline(False);expect(page.get_by_text(re.compile('^待上传课记与选项：'))).to_have_count(0,timeout=20000);assert len(state()['lessons'])==count+1
    click('我的账号');shot('account','我的账号与数据导出','账号与初始设置','数据导出、导入、登录设备和本机显示集中在账号，避免占据今天页。')
    with page.expect_download() as download:click('导出 Excel')
    file=OUT/'synthetic-export.xlsx';download.value.save_as(file);book=load_workbook(file,data_only=False);assert not any(c.data_type=='f' for sheet in book for row in sheet for c in row)
    checks.append('实际浏览器下载XLSX可读取，导出无可执行公式；没有保存真实业务文件')
    shot('export','Excel实际下载确认','账号与初始设置','本次点击实际产生Excel下载并重新读取检查；本机未提交草稿、待上传和登录凭据不导出。',page.get_by_text('Excel 已交给浏览器下载，请在下载列表中查看。',exact=True))
    click('使用说明与手机桌面入口');shot('help','账号角落的问号说明','账号与初始设置','帮助集中在账号问号，保留方法证据边界、家庭帮助说明与手机桌面入口。');close_all()
    click('维护课程和老师选项');shot('catalog','课程与老师选项维护','账号与初始设置','新增、改名、停用和恢复由这里维护；历史课程保留当时名称。')
    click('新增领域');expect(btn('保存')).to_be_disabled();click('放弃这份填写');shot('discard','放弃本机填写需确认','观察与可靠性','空名称不能保存；放弃只清除本机未提交输入，已保存数据不变，结果未知请求不提供直接丢弃。',btn('确认放弃这份填写'));click('确认放弃这份填写')
    click('我的账号');click('打开导入工具')
    w=Workbook();ws=w.active;ws.title='演示目标';ws.append(['编号','名称']);ws.append(['AUDIT-G1','演示：导入后核对的目标']);upload=OUT/'synthetic-import.xlsx';w.save(upload)
    with page.expect_file_chooser() as chooser:click('选择 Excel／CSV')
    chooser.value.set_files(upload);click('上传并读取工作表');select('演示目标：用途','goal');select('演示目标：来源编号（必选，原表内唯一）','编号');select('演示目标：名称／内容','名称');field('资料来源标识').fill('synthetic-review-20261029');click('核对映射并生成写入预览')
    expect(page.get_by_text('写入预览 · 1 行',exact=True)).to_be_visible();shot('import','Excel显式映射与写入预览','账号与初始设置','本次只到预览，不自动写入。实际导入、重复去重、原文来源与回执重试另有完整合约及浏览器验收。',page.get_by_text('写入预览 · 1 行',exact=True))
    checks.append('截图流程覆盖真实写入、关联、共享计划、课程整理、反馈、复盘、测量、冲突、离线补传和文件下载；所有页面JS异常为0')
    assert not errors,errors
    (OUT/'manifest.json').write_text(json.dumps({'source':source,'synthetic':True,'shots':shots,'checks':checks},ensure_ascii=False,indent=2)+'\n')
    print(f'PASS curated review journey: {len(shots)} screenshots, real persistence assertions, zero browser exceptions',flush=True)
    browser.close()
