"""Real browser multipart/mapping/retry/source roundtrip; only generated synthetic data."""
from prd_ui_helpers import reveal_control
import os,re
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import Workbook,load_workbook
from playwright.sync_api import sync_playwright,expect
BASE=os.environ['SKARMA_TEST_URL'].rstrip('/');assert urlparse(BASE).hostname in {'localhost','127.0.0.1','::1'}
OUT=Path('.local/prd-import-ui');OUT.mkdir(parents=True,exist_ok=True)
w=Workbook();s=w.active;s.title='目标';s.append(['编号','名称','未映射']);s.append(['G1','合成导入目标','=SUM(A1:A3)']);s=w.create_sheet('方法');s.append(['编号','名称','关联目标']);s.append(['M1','合成导入方法','目标:G1']);s=w.create_sheet('经验');s.append(['编号','内容','关联目标','阶段','状态']);s.append(['K1','合成既有经验','目标:G1','已验证','停用']);s=w.create_sheet('历史');s.append(['原文']);s.append(['合成原始历史🙂']);filename=OUT/'synthetic.xlsx';w.save(filename)
with sync_playwright() as p:
 b=p.chromium.launch(executable_path=os.environ.get('CHROMIUM_EXECUTABLE'),headless=True,args=['--no-sandbox']);ctx=b.new_context(viewport={'width':360,'height':740},accept_downloads=True);page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
 def button(name):return page.get_by_role('button',name=name,exact=True)
 def click(name):reveal_control(page,'button',name).click()
 def state():return ctx.request.get(BASE+'/api/prd/state').json()
 page.goto(BASE);reveal_control(page,'textbox','账号').fill('owner');page.get_by_label('密码',exact=True).fill('synthetic-prd-ui-password');click('登录');expect(button('记录上课')).to_be_enabled();click('我的账号');click('打开导入工具')
 def choose():
  with page.expect_file_chooser() as chooser:button(re.compile('^(重新)?选择 Excel／CSV$')).click()
  chooser.value.set_files(filename);click('上传并读取工作表');expect(page.get_by_role('combobox',name='目标：用途',exact=True)).to_be_visible()
 def map_rules():
  reveal_control(page,'textbox','资料来源标识').fill('synthetic-browser-source')
  for sheet,kind in [('目标','goal'),('方法','method'),('经验','experience'),('历史','archive')]:page.get_by_role('combobox',name=sheet+'：用途',exact=True).select_option(kind)
  for sheet in ['目标','方法']:
   page.get_by_role('combobox',name=sheet+'：来源编号（必选，原表内唯一）',exact=True).select_option('编号');page.get_by_role('combobox',name=sheet+'：名称／内容',exact=True).select_option('名称')
  reveal_control(page,'combobox','方法：明确关联目标来源').select_option('关联目标');
  for key,column in [('来源编号（必选，原表内唯一）','编号'),('名称／内容','内容'),('明确关联目标来源','关联目标'),('明确经验阶段（提出／已使用／已验证）','阶段'),('经验使用状态（启用／暂停／停用）','状态')]:page.get_by_role('combobox',name='经验：'+key,exact=True).select_option(column)
  click('核对映射并生成写入预览');expect(page.get_by_text('写入预览 · 4 行',exact=True)).to_be_visible()
 choose();map_rules();assert not state()['goals'];page.get_by_role('checkbox',name=re.compile('我已核对用途')).check()
 # The server commits but the response vanishes. An exact retry must return the same receipt.
 def lost(route):
  response=route.fetch();assert response.ok,response.text();route.abort('failed')
 page.route('**/api/data-transfer',lambda route:lost(route) if route.request.post_data_json['kind']=='commit' else route.continue_(),times=1)
 click('确认按此预览导入');expect(page.get_by_text(re.compile('暂时无法连接'))).to_be_visible();assert len(state()['goals'])==1;click('确认按此预览导入');expect(page.get_by_text('导入回执',exact=True)).to_be_visible();assert len(state()['goals'])==1;assert len(state()['methods'])==1;assert state()['knowledge'][0]['input']['stage']=='helpful';assert state()['knowledge'][0]['details']['state']=='retired'
 button('核对来源：目标 第2行').click();expect(page.get_by_text('完整来源资料',exact=True)).to_be_visible();assert '=SUM(A1:A3)' in page.get_by_text(re.compile('"batch_id"')).last.inner_text();click('收起来源资料')
 choose();map_rules();expect(page.get_by_text('预计新增业务 0，仅存档 0，相同来源跳过 4。',exact=True)).to_be_visible();page.get_by_role('checkbox',name=re.compile('我已核对用途')).check();click('确认按此预览导入');expect(page.get_by_text(re.compile('服务器已确认导入'))).to_be_visible();assert len(state()['goals'])==1
 click('查看最近导入批次');expect(button(re.compile('^查看批次 '))).to_have_count(2)
 with page.expect_download() as download:click('导出 Excel')
 target=OUT/'synthetic-export.xlsx';download.value.save_as(target);book=load_workbook(target,data_only=False);assert '导入来源' in book.sheetnames and '导入批次' in book.sheetnames
 assert any(cell.value=='=SUM(A1:A3)' or '=SUM(A1:A3)' in str(cell.value) for row in book['导入来源'] for cell in row);assert not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row)
 assert page.evaluate('document.documentElement.scrollWidth <= innerWidth+1');assert not errors,errors
 page.screenshot(path=str(OUT/'import-mobile.png'),full_page=False);print('PASS browser explicit mapping, target-method relation, lost-response replay, duplicate skip, source audit and provenance Excel',flush=True);b.close()
