"""Download new-model workbook from an actual browser, excluding unsubmitted local inputs."""
from prd_ui_helpers import reveal_control
import os,json,re
from pathlib import Path
from urllib.parse import urlparse
from uuid import uuid4
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright,expect
BASE=os.environ['SKARMA_TEST_URL'].rstrip('/');assert urlparse(BASE).hostname in {'localhost','127.0.0.1','::1'}
OUT=Path('.local/prd-export-ui');OUT.mkdir(parents=True,exist_ok=True)
with sync_playwright() as p:
 b=p.chromium.launch(executable_path=os.environ.get('CHROMIUM_EXECUTABLE'),headless=True,args=['--no-sandbox']);ctx=b.new_context(viewport={'width':360,'height':740},accept_downloads=True);page=ctx.new_page();errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
 def btn(name):return page.get_by_role('button',name=name,exact=True)
 def click(name):reveal_control(page,'button',name).click()
 def field(name):return reveal_control(page,'textbox',name)
 def call(method,path,body=None):
  r=ctx.request.fetch(BASE+'/api'+path,method=method,data=body,headers={'X-Skarma-Request':'1','X-Skarma-Contract':'skarmar.prd.v1'});assert r.ok,r.text();return r.json()
 page.goto(BASE);field('账号').fill('owner');page.get_by_label('密码',exact=True).fill('synthetic-prd-ui-password');click('登录');expect(btn('记录课堂')).to_be_enabled()
 domain_id=str(uuid4());domain_saved=call('PUT','/prd/records/'+domain_id,{'id':domain_id,'operation_id':str(uuid4()),'expected_version':None,'change_note':'','mutation':{'kind':'catalog','value':{'kind':'domain','name':'测试领域','active':True}}})['value']
 domain={'id':domain_saved['id'],'version':domain_saved['version'],'kind':'domain','name':domain_saved['input']['name']}
 goal=str(uuid4());call('PUT','/prd/records/'+goal,{'id':goal,'operation_id':str(uuid4()),'expected_version':None,'change_note':'','mutation':{'kind':'goal','value':{'title':'=SUM(A1:A9)','level':'short','meaning':'合成业务记录','domain':domain,'due':'2027-01-20','support':'先不提示，需要时再示范'}}})
 options={}
 for kind in ['course','teacher']:
  id=str(uuid4());options[kind]=id;call('PUT','/prd/records/'+id,{'id':id,'operation_id':str(uuid4()),'expected_version':None,'change_note':'','mutation':{'kind':'catalog','value':{'kind':kind,'name':'测试'+kind,'active':True}}})
 page.reload();expect(btn('记录课堂')).to_be_enabled();click('目标');click('新建目标');field('本领名称').fill('LOCAL_UNSUBMITTED_GOAL_9A');click('关闭填写目标');click('今天');click('记录课堂')
 for prefix,label,kind in [('课程名称','课程','course'),('授课老师','老师','teacher')]:
  click(prefix+'（必填）：请选择'+label);page.get_by_role('dialog',name='选择'+label,exact=True).get_by_role('button',name='测试'+kind,exact=True).click()
 click('补充观察');field('今天课上怎么样？').fill('LOCAL_QUEUED_LESSON_9A');page.route('**/api/prd/queue',lambda r:r.abort('failed'));click('提交记录');expect(page.get_by_role('dialog')).to_have_count(0);expect(btn('检查连接并继续上传')).to_be_visible()
 before=call('GET','/prd/state');assert not before['lessons'];click('我的账号');expect(btn('导出 Excel')).to_be_visible()
 with page.expect_download() as event:click('导出 Excel')
 download=event.value;assert download.suggested_filename.startswith('skarmar-');target=OUT/'synthetic-new-model.xlsx';download.save_as(target);book=load_workbook(target,data_only=False)
 assert book['导出说明']['B1'].value=='skarmar.prd.export.v1';assert '实际测量' in book.sheetnames and '复盘修订' in book.sheetnames
 texts='\n'.join(str(cell.value) for sheet in book for row in sheet for cell in row if cell.value is not None)
 assert 'LOCAL_UNSUBMITTED_GOAL_9A' not in texts;assert 'LOCAL_QUEUED_LESSON_9A' not in texts;assert 'synthetic-prd-ui-password' not in texts
 goal_headers=[cell.value for cell in book['目标'][1]]
 assert {'填写 · 领域','填写 · 目标日期','填写 · 帮助条件'}.issubset(set(goal_headers))
 assert '测试领域' in texts and '2027-01-20' in texts and '先不提示，需要时再示范' in texts
 assert any(cell.value=='=SUM(A1:A9)' and cell.data_type=='s' for row in book['目标'] for cell in row);assert not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row)
 assert call('GET','/prd/state')==before
 # Static mobile metadata/assets are served by the application, with the product spelling.
 manifest=ctx.request.get(BASE+'/manifest.webmanifest');assert manifest.ok and manifest.json()['name']=='skarmar'
 for path in ['/icon-192.png','/icon-512.png']:
  r=ctx.request.get(BASE+path);assert r.ok and r.body().startswith(b'\x89PNG')
 assert page.locator('html').get_attribute('lang')=='zh-CN';assert not errors,errors
 print('PASS browser XLSX download, excludes local draft/queue/password, literal formula text, unchanged server data and install assets',flush=True);b.close()
